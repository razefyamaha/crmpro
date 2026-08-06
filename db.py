"""
db.py - Excel database layer for CRM Pro.

Reads and writes data/CRM_Database.xlsx using openpyxl.
The Excel file IS the database: any edits made in the app are saved
straight back into the workbook, and edits you make in Excel are picked
up when the app reloads.
"""
import os
import sys
import threading
from datetime import datetime
from openpyxl import load_workbook

_lock = threading.Lock()


def _app_data_dir():
    # When frozen into an .exe, keep the data folder next to the .exe so the
    # Excel database is user-visible and persists. When run from source, use
    # the project's data folder.
    if getattr(sys, "frozen", False):
        base = os.path.dirname(os.path.abspath(sys.executable))
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, "data")


DATA_DIR = _app_data_dir()
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(DATA_DIR, "CRM_Database.xlsx")

# Column indexes (1-based) in the Quotations sheet
Q = {"QNO": 1, "DATE": 2, "TIME": 3, "COMPANY": 4, "STATUS": 5, "AMOUNT": 6, "BY": 7, "LINES": 8}
# Companies sheet
C = {"COMPANY": 1, "INDUSTRY": 2, "CONTACT": 3, "EMAIL": 4, "PHONE": 5,
     "SOURCE": 6, "STAGE": 7, "NOTES": 8, "ADDRESS": 9, "PIPELINE": 10}

# Lead pipeline stages (in order) tracked per customer with dates
PIPELINE_STAGES = ["Sent Email", "Contacted", "Sent Quotation", "In Progress"]

STATUSES = ["New", "Quotation In Progress", "Completed", "Deal Won", "Deal Lost"]
STAGES = ["New", "Contacted", "Quoted", "Negotiating", "Won", "Lost"]


def _load_sheet(name):
    wb = load_workbook(DB_PATH, data_only=True)
    ws = wb[name]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in r):
            continue
        rows.append(r)
    wb.close()
    return rows


def _load_lines(raw):
    """Parse the stored line-items JSON into a list of dicts."""
    if not raw:
        return []
    try:
        import json
        data = json.loads(raw)
        if isinstance(data, list):
            return data
    except Exception:
        pass
    return []


def _write_quotations(rows):
    """rows: list of lists (8 cols) in column order."""
    wb = load_workbook(DB_PATH)
    ws = wb["Quotations"]
    # clear data rows
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    for i, row in enumerate(rows):
        ws.append(row)
    wb.save(DB_PATH)
    wb.close()


def _write_companies(rows):
    wb = load_workbook(DB_PATH)
    ws = wb["Companies"]
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    for row in rows:
        ws.append(row)
    wb.save(DB_PATH)
    wb.close()


# ---------------- Quotations ----------------
def load_quotations():
    with _lock:
        data = _load_sheet("Quotations")
    out = []
    for i, row in enumerate(data):
        amount = row[Q["AMOUNT"] - 1] or 0
        try:
            amount = float(amount)
        except (TypeError, ValueError):
            amount = 0.0
        date = row[Q["DATE"] - 1] or ""
        if hasattr(date, "strftime"):
            date = date.strftime("%Y-%m-%d")
        out.append({
            "id": i,
            "qno": row[Q["QNO"] - 1],
            "date": str(date),
            "time": row[Q["TIME"] - 1] or "",
            "company": row[Q["COMPANY"] - 1],
            "status": row[Q["STATUS"] - 1],
            "amount": amount,
            "created_by": row[Q["BY"] - 1],
            "lines": _load_lines(row[Q["LINES"] - 1] if len(row) >= Q["LINES"] else None),
        })
    return out


def save_quotations(quotes):
    """quotes: list of dicts (must contain all keys)."""
    import json
    with _lock:
        rows = [[q["qno"], q["date"], q["time"], q["company"], q["status"],
                 q["amount"], q["created_by"], json.dumps(q.get("lines", []))]
                for q in quotes]
        _write_quotations(rows)


def next_qno(quotes):
    nums = []
    for q in quotes:
        m = str(q.get("qno") or "")
        if m.upper().startswith("QT") and m[2:].isdigit():
            nums.append(int(m[2:]))
    nxt = (max(nums) + 1) if nums else 1
    return f"QT{nxt:05d}"


# ---------------- Companies ----------------
def load_companies():
    with _lock:
        data = _load_sheet("Companies")
    out = []
    for i, row in enumerate(data):
        out.append({
            "id": i,
            "company": row[C["COMPANY"] - 1] or "",
            "industry": row[C["INDUSTRY"] - 1] or "",
            "contact": row[C["CONTACT"] - 1] or "",
            "email": row[C["EMAIL"] - 1] or "",
            "phone": row[C["PHONE"] - 1] or "",
            "source": row[C["SOURCE"] - 1] or "",
            "stage": row[C["STAGE"] - 1] or "New",
            "notes": row[C["NOTES"] - 1] or "",
            "address": row[C["ADDRESS"] - 1] if len(row) >= C["ADDRESS"] else "",
            "pipeline": _load_pipeline(row[C["PIPELINE"] - 1] if len(row) >= C["PIPELINE"] else None),
        })
    return out


def _load_pipeline(raw):
    """Parse stored pipeline JSON into a dict of stage->date."""
    default = {s: "" for s in PIPELINE_STAGES}
    if not raw:
        return default
    try:
        import json
        data = json.loads(raw)
        if isinstance(data, dict):
            for k in default:
                if k in data:
                    default[k] = str(data[k] or "")
            return default
    except Exception:
        pass
    return default


def save_pipeline(company, stage, date_val):
    """Record a date for a pipeline stage on the given company."""
    companies = load_companies()
    for c in companies:
        if c["company"].lower() == company.lower():
            c["pipeline"][stage] = str(date_val or "")
            break
    save_companies(companies)


def save_companies(companies):
    import json
    with _lock:
        rows = [[c["company"], c["industry"], c["contact"], c["email"],
                 c["phone"], c["source"], c["stage"], c["notes"], c.get("address", ""),
                 json.dumps(c.get("pipeline", {}))]
                for c in companies]
        _write_companies(rows)


# Products sheet
P = {"ID": 1, "NAME": 2, "CATEGORY": 3, "COST": 4, "SRP": 5, "MARGIN": 6,
     "SELLING": 7, "PROFIT": 8, "NOTES": 9}
# Settings sheet
S = {"KEY": 1, "VALUE": 2}

# Users sheet
U = {"USERNAME": 1, "PASSWORD": 2, "NAME": 3, "ROLE": 4, "ACTIVE": 5}

ROLES = ["superadmin", "admin", "user"]
# Forecasts sheet
F = {"NO": 1, "DATE": 2, "CUSTOMER": 3, "VALUE": 4, "COST": 5, "PROFIT": 6,
     "ITEMS": 7, "STATUS": 8, "NOTES": 9, "LINES": 10}
# Leads (contacts) sheet - standalone, separate from customers
L = {"NAME": 1, "COMPANY": 2, "PHONE": 3, "EMAIL": 4, "SOURCE": 5, "STATUS": 6,
     "DATE": 7, "NOTES": 8}

# Lead statuses; 'Quotation In Progress' moves the lead into the quotation list
LEAD_STATUSES = ["New", "Contacted", "Quotation In Progress", "Deal Won", "Deal Lost"]


def _ensure_sheets():
    """Create Products + Settings + Forecasts sheets if missing (idempotent)."""
    from openpyxl import load_workbook
    wb = load_workbook(DB_PATH)
    changed = False
    if "Products" not in wb.sheetnames:
        ws = wb.create_sheet("Products")
        ws.append(["Product ID", "Product Name", "Category", "Cost Price (RM)",
                   "SRP (RM)", "Margin %", "Selling Price (RM)", "Profit (RM)", "Notes"])
        for c in ws.iter_cols(1, 9):
            ws.column_dimensions[c[0].column_letter].width = 18
        changed = True
    if "Settings" not in wb.sheetnames:
        ws = wb.create_sheet("Settings")
        ws.append(["Key", "Value"])
        ws.append(["monthly_target", 60000])
        changed = True
    if "Forecasts" not in wb.sheetnames:
        ws = wb.create_sheet("Forecasts")
        ws.append(["Forecast No", "Date", "Customer", "Total Value (RM)",
                   "Total Cost (RM)", "Total Profit (RM)", "Items", "Status",
                   "Notes", "Line Items (JSON)"])
        for c in ws.iter_cols(1, 10):
            ws.column_dimensions[c[0].column_letter].width = 18
        changed = True
    if "Leads" not in wb.sheetnames:
        ws = wb.create_sheet("Leads")
        ws.append(["Name", "Company", "Phone", "Email", "Lead Source", "Status",
                   "Date", "Notes"])
        for c in ws.iter_cols(1, 8):
            ws.column_dimensions[c[0].column_letter].width = 20
        changed = True
    if "Users" not in wb.sheetnames:
        ws = wb.create_sheet("Users")
        ws.append(["Username", "Password Hash", "Display Name", "Role", "Active"])
        for c in ws.iter_cols(1, 5):
            ws.column_dimensions[c[0].column_letter].width = 20
        # Seed the default super admin account (username: admin / password: admin123)
        from werkzeug.security import generate_password_hash
        ws.append(["admin", generate_password_hash("admin123"), "Super Admin", "superadmin", 1])
        changed = True
    else:
        # ensure the Line Items column header exists
        ws = wb["Forecasts"]
        hdr = [ws.cell(row=1, column=cc).value for cc in range(1, ws.max_column + 1)]
        if "Line Items" not in hdr:
            ws.cell(row=1, column=F["LINES"], value="Line Items (JSON)")
            changed = True
    # ensure Quotations sheet has a Line Items column header
    if "Quotations" in wb.sheetnames:
        ws = wb["Quotations"]
        hdr = [ws.cell(row=1, column=cc).value for cc in range(1, ws.max_column + 1)]
        if "Line Items" not in hdr:
            ws.cell(row=1, column=Q["LINES"], value="Line Items (JSON)")
            changed = True
    # ensure Companies sheet has Address + Pipeline column headers
    if "Companies" in wb.sheetnames:
        ws = wb["Companies"]
        hdr = [ws.cell(row=1, column=cc).value for cc in range(1, ws.max_column + 1)]
        if "Address" not in hdr:
            ws.cell(row=1, column=C["ADDRESS"], value="Address")
            changed = True
        if "Pipeline" not in hdr:
            ws.cell(row=1, column=C["PIPELINE"], value="Pipeline (JSON)")
            changed = True
    if changed:
        wb.save(DB_PATH)
    wb.close()


# ---------------- Forecasts ----------------
def load_forecasts():
    _ensure_sheets()
    with _lock:
        data = _load_sheet("Forecasts")
    out = []
    for i, row in enumerate(data):
        def num(v):
            try:
                return float(v) if v not in (None, "") else 0.0
            except (TypeError, ValueError):
                return 0.0
        out.append({
            "id": i,
            "fc_no": str(row[F["NO"] - 1] or ""),
            "date": str(row[F["DATE"] - 1] or ""),
            "customer": str(row[F["CUSTOMER"] - 1] or ""),
            "value": num(row[F["VALUE"] - 1]),
            "cost": num(row[F["COST"] - 1]),
            "profit": num(row[F["PROFIT"] - 1]),
            "items": str(row[F["ITEMS"] - 1] or ""),
            "status": str(row[F["STATUS"] - 1] or "Open"),
            "notes": str(row[F["NOTES"] - 1] or ""),
            "lines": _load_lines(row[F["LINES"] - 1] if len(row) >= F["LINES"] else None),
        })
    return out


def _load_lines(raw):
    """Parse the stored line-items JSON into a list of dicts."""
    if not raw:
        return []
    try:
        import json
        data = json.loads(raw)
        if isinstance(data, list):
            return data
    except Exception:
        pass
    return []


def save_forecasts(forecasts):
    _ensure_sheets()
    import json
    with _lock:
        wb = load_workbook(DB_PATH)
        ws = wb["Forecasts"]
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        for f in forecasts:
            lines_json = json.dumps(f.get("lines", []))
            ws.append([f["fc_no"], f["date"], f["customer"], f["value"], f["cost"],
                       f["profit"], f["items"], f["status"], f["notes"], lines_json])
        wb.save(DB_PATH)
        wb.close()


def next_forecast_no(forecasts):
    nums = []
    for f in forecasts:
        n = str(f.get("fc_no") or "")
        if n.upper().startswith("FC") and n[2:].isdigit():
            nums.append(int(n[2:]))
    nxt = (max(nums) + 1) if nums else 1
    return f"FC{nxt:04d}"


# ---------------- Leads (contacts) ----------------
def load_leads():
    _ensure_sheets()
    with _lock:
        data = _load_sheet("Leads")
    out = []
    for i, row in enumerate(data):
        out.append({
            "id": i,
            "name": str(row[L["NAME"] - 1] or ""),
            "company": str(row[L["COMPANY"] - 1] or ""),
            "phone": str(row[L["PHONE"] - 1] or ""),
            "email": str(row[L["EMAIL"] - 1] or ""),
            "source": str(row[L["SOURCE"] - 1] or ""),
            "status": str(row[L["STATUS"] - 1] or "New"),
            "date": str(row[L["DATE"] - 1] or ""),
            "notes": str(row[L["NOTES"] - 1] or ""),
        })
    return out


def save_leads(leads):
    _ensure_sheets()
    with _lock:
        wb = load_workbook(DB_PATH)
        ws = wb["Leads"]
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        for ld in leads:
            ws.append([ld["name"], ld["company"], ld["phone"], ld["email"],
                       ld["source"], ld["status"], ld["date"], ld["notes"]])
        wb.save(DB_PATH)
        wb.close()


# ---------------- Products ----------------
def load_products():
    _ensure_sheets()
    with _lock:
        data = _load_sheet("Products")
    out = []
    for i, row in enumerate(data):
        def num(v):
            try:
                return float(v) if v not in (None, "") else 0.0
            except (TypeError, ValueError):
                return 0.0
        out.append({
            "id": i,
            "pid": str(row[P["ID"] - 1] or ""),
            "name": str(row[P["NAME"] - 1] or ""),
            "category": str(row[P["CATEGORY"] - 1] or ""),
            "cost": num(row[P["COST"] - 1]),
            "srp": num(row[P["SRP"] - 1]),
            "margin": num(row[P["MARGIN"] - 1]),
            "selling": num(row[P["SELLING"] - 1]),
            "profit": num(row[P["PROFIT"] - 1]),
            "notes": str(row[P["NOTES"] - 1] or ""),
        })
    return out


def save_products(products):
    _ensure_sheets()
    with _lock:
        wb = load_workbook(DB_PATH)
        ws = wb["Products"]
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        for p in products:
            ws.append([p["pid"], p["name"], p["category"], p["cost"], p["srp"],
                       p["margin"], p["selling"], p["profit"], p["notes"]])
        wb.save(DB_PATH)
        wb.close()


def next_product_id(products):
    nums = []
    for p in products:
        pid = str(p.get("pid") or "")
        if pid.upper().startswith("PRD") and pid[3:].isdigit():
            nums.append(int(pid[3:]))
    nxt = (max(nums) + 1) if nums else 1
    return f"PRD{nxt:03d}"


# ---------------- Users ----------------
def load_users():
    """Return a list of user dicts from the Users sheet."""
    _ensure_sheets()
    with _lock:
        rows = _load_sheet("Users")
    users = []
    for row in rows:
        if not row[U["USERNAME"] - 1]:
            continue
        users.append({
            "username": str(row[U["USERNAME"] - 1]),
            "password": str(row[U["PASSWORD"] - 1] or ""),
            "name": str(row[U["NAME"] - 1] or ""),
            "role": str(row[U["ROLE"] - 1] or "user"),
            "active": 1 if row[U["ACTIVE"] - 1] in (1, True, "1", "TRUE", "true") else 0,
        })
    return users


def save_users(users):
    _ensure_sheets()
    with _lock:
        wb = load_workbook(DB_PATH)
        ws = wb["Users"]
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        for u in users:
            ws.append([u["username"], u.get("password", ""), u.get("name", ""),
                       u.get("role", "user"), 1 if u.get("active") else 0])
        wb.save(DB_PATH)
        wb.close()


def find_user(username):
    username = (username or "").strip()
    for u in load_users():
        if u["username"].lower() == username.lower():
            return u
    return None


# ---------------- Settings / Targets ----------------
def get_setting(key, default=None):
    _ensure_sheets()
    with _lock:
        data = _load_sheet("Settings")
    for row in data:
        if str(row[S["KEY"] - 1] or "") == key:
            return row[S["VALUE"] - 1]
    return default


def set_setting(key, value):
    _ensure_sheets()
    with _lock:
        wb = load_workbook(DB_PATH)
        ws = wb["Settings"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        found = False
        for r, row in enumerate(rows, start=2):
            if str(row[S["KEY"] - 1] or "") == key:
                ws.cell(row=r, column=S["VALUE"], value=value)
                found = True
        if not found:
            ws.append([key, value])
        wb.save(DB_PATH)
        wb.close()


def monthly_target():
    try:
        return float(get_setting("monthly_target", 60000))
    except (TypeError, ValueError):
        return 60000.0


def current_theme():
    """Return the currently selected UI theme (one of blue/emerald/sunset)."""
    t = get_setting("ui_theme", "blue")
    if t not in ("blue", "emerald", "sunset"):
        return "blue"
    return t


def today():
    return datetime.now().strftime("%Y-%m-%d")


def now_time():
    return datetime.now().strftime("%I:%M %p").lstrip("0")


# ---------------- Uploads (customer document folders) ----------------
def uploads_root():
    return os.path.join(DATA_DIR, "uploads")


def customer_folder(company):
    """Return (and create) the folder that holds a customer's uploaded files."""
    import re as _re
    slug = _re.sub(r"[^\w\- ]+", "", str(company or "")).strip().replace(" ", "_")
    slug = slug or "uncategorized"
    folder = os.path.join(uploads_root(), slug)
    os.makedirs(folder, exist_ok=True)
    return folder


def customer_file_path(company, fname):
    """Full path to a customer's uploaded file (even if it does not exist)."""
    return os.path.join(customer_folder(company), str(fname or ""))


def list_uploads(company):
    """Return a list of {name, size, mtime} for a customer's uploaded files."""
    folder = customer_folder(company)
    files = []
    try:
        for name in os.listdir(folder):
            path = os.path.join(folder, name)
            if os.path.isfile(path):
                st = os.stat(path)
                files.append({"name": name, "size": st.st_size,
                              "mtime": datetime.fromtimestamp(st.st_mtime)
                              .strftime("%d %b %Y %I:%M %p").lstrip("0")})
    except Exception:
        pass
    files.sort(key=lambda f: f["mtime"], reverse=True)
    return files


def thumbnail_path(company, fname):
    """Path to the cached first-page thumbnail for a PDF (None if not a PDF)."""
    if not fname.lower().endswith(".pdf"):
        return None
    folder = customer_folder(company)
    thumb_dir = os.path.join(folder, "_thumbs")
    os.makedirs(thumb_dir, exist_ok=True)
    return os.path.join(thumb_dir, fname + ".jpg")


def render_thumbnail(pdf_path, thumb_path, width=360):
    """Render the first page of a PDF to a cached JPEG thumbnail."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(pdf_path)
    try:
        if len(pdf) == 0:
            return None
        page = pdf[0]
        bitmap = page.render(scale=0.5)
        img = bitmap.to_pil()
        img.thumbnail((width, width))
        img.save(thumb_path, "JPEG", quality=82)
        return thumb_path
    finally:
        pdf.close()


def pdf_page_count(pdf_path):
    """Return the number of pages in a PDF (1 if unknown/error)."""
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(pdf_path)
        try:
            return max(len(pdf), 1)
        finally:
            pdf.close()
    except Exception:
        return 1


def page_image_path(company, fname, page_num):
    """Cached path to a rendered image of a specific PDF page."""
    folder = customer_folder(company)
    pdir = os.path.join(folder, "_pages")
    os.makedirs(pdir, exist_ok=True)
    return os.path.join(pdir, f"{fname}__p{page_num}.jpg")


def render_page(pdf_path, page_num, out_path, scale=1.7):
    """Render one page of a PDF to a cached JPEG image (0-based page index)."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(pdf_path)
    try:
        total = len(pdf)
        if page_num < 0 or page_num >= total:
            return None
        page = pdf[page_num]
        bitmap = page.render(scale=scale)
        img = bitmap.to_pil()
        img.convert("RGB").save(out_path, "JPEG", quality=88)
        return out_path
    finally:
        pdf.close()
